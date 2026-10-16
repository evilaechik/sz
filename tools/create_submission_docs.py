from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('docs/submission')

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), fill)
    tc_pr.append(shading)

def set_cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    mar = tc_pr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar')
        tc_pr.append(mar)
    for side, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None:
            node = OxmlElement(f'w:{side}')
            mar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def configure(doc):
    section = doc.sections[0]
    section.top_margin = Inches(.72)
    section.bottom_margin = Inches(.72)
    section.left_margin = Inches(.78)
    section.right_margin = Inches(.78)
    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12)]:
        style = doc.styles[name]
        style.font.name = 'Aptos Display'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
    # Keep headers empty: explicit page breaks and long tables render more
    # consistently across Word, LibreOffice and competition PDF viewers.
    section.header.paragraphs[0].text = ''

def title(doc, text, subtitle):
    p = doc.add_paragraph(style='Title')
    p.add_run(text)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(subtitle)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(91, 109, 228)

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = 'Table Grid'
    for i, header in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, '183054')
        set_cell_margin(cell)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)
    header_props = t.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement('w:tblHeader')
    repeat_header.set(qn('w:val'), 'true')
    header_props.append(repeat_header)
    for row_i, row in enumerate(rows):
        cells = t.add_row().cells
        row_props = t.rows[-1]._tr.get_or_add_trPr()
        row_props.append(OxmlElement('w:cantSplit'))
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margin(cells[i])
            if row_i % 2 == 1:
                set_cell_shading(cells[i], 'F4F7FC')
            for p in cells[i].paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8.8)
        if widths:
            for i, width in enumerate(widths):
                cells[i].width = Inches(width)
    doc.add_paragraph()
    return t

def bullet(doc, text):
    doc.add_paragraph(text, style='List Bullet')

def technical_doc():
    doc = Document(); configure(doc)
    title(doc, 'Runway Guard Technical Documentation', 'Topic E  SME Finance Copilot  |  Submission version 1.1')
    doc.add_heading('1. Product purpose', level=1)
    doc.add_paragraph('Runway Guard is an AI finance copilot for small-business owners. It turns normalized banking, payments and accounting activity into a cash-flow assessment, reconciliation review, safe collection-reminder draft and approval-ready recommendations. The prototype uses synthetic Harbor & Bean data and simulated financial operations only.')
    doc.add_heading('2. User problem and task completion', level=1)
    doc.add_paragraph('Small-business owners often discover liquidity risk too late and lack a single, safe place to understand transactions, receivables and actions. Runway Guard prioritizes a projected Week 43 cash shortfall, identifies the overdue Northstar receivable, and prepares interventions without independently moving funds or sending customer messages.')
    table(doc, ['Task', 'Agent behavior', 'Human control'], [
        ['Cash assessment', 'Reads normalized transactions and reports inflows, outflows and forecast risk.', 'Read-only analysis.'],
        ['Reconciliation', 'Finds rows requiring review and prepares evidence-backed suggestions.', 'No ledger posting.'],
        ['Collections', 'Drafts a reminder for a known overdue invoice.', 'Owner approval before delivery.'],
        ['Security review', 'Classifies unverified bank-detail changes and policy-override attempts.', 'Unsafe instruction is blocked.'],
        ['Sandbox evidence', 'Runs simulated Alipay or WeChat Pay checks and exports logs.', 'No production key, customer data or payment movement.'],
    ], [1.2, 3.65, 1.7])
    doc.add_page_break()
    doc.add_heading('3. System architecture', level=1)
    doc.add_paragraph('The React client renders the owner workspace, approval queue, integration state and trace. It calls a Node/Express API. The API owns tool execution, policy checks, sandbox evidence and the optional model boundary. The browser never receives an OpenAI key or a payment credential.')
    table(doc, ['Layer', 'Responsibility', 'Boundary'], [
        ['Client', 'Dashboard, CSV import, prompts, approvals and readable traces.', 'Calls only local API endpoints.'],
        ['Agent API', 'Routes requests, executes allowlisted tools and returns structured tool evidence.', 'Rejects unknown tools and malformed requests.'],
        ['Optional LLM planner', 'Uses the OpenAI Responses API when a server-side key is configured.', 'May plan and call tools, but cannot bypass policy.'],
        ['Policy layer', 'Requires approval for consequential drafts and blocks unsafe settlement instructions.', 'No payment, message or ledger-posting tool exists.'],
        ['Sandbox evidence', 'Produces simulated Alipay and WeChat Pay operation logs.', 'Explicitly non-production.'],
    ], [1.2, 3.65, 1.7])
    doc.add_heading('4. Core algorithm and agent loop', level=1)
    doc.add_paragraph('The API first validates the prompt and a bounded transaction payload. With an OpenAI key, the planner receives a narrow tool schema and may request an allowlisted function. The server executes the function against normalized data, returns structured output to the planner, and limits the loop to three turns. Without a key, a deterministic local policy agent executes the same safe tool paths for a reproducible demonstration.')
    bullet(doc, 'Normalize data from CSV and demo adapters into date, merchant, amount, category and review-status fields.')
    bullet(doc, 'Select the relevant read-only or draft-only tool based on the request.')
    bullet(doc, 'Return plan, tool calls, outcome and approval requirement to the owner.')
    bullet(doc, 'Record sandbox execution evidence separately from business recommendations.')
    doc.add_heading('5. Data and integrations', level=1)
    doc.add_paragraph('The demo includes synthetic banking, Stripe-like payments and accounting records, plus consent-gated Alipay and WeChat Pay sandbox adapters. CSV input is parsed in the browser and valid rows are marked for review. Sandbox tests simulate merchant-profile activation, read-only settlement metadata retrieval and signed callback verification. They do not accept production credentials or connect to external merchant accounts.')
    doc.add_page_break()
    doc.add_heading('6. Deployment and test instructions', level=1)
    table(doc, ['Step', 'Command or action'], [
        ['Prerequisite', 'Node.js 20.19 or newer and npm'], ['Install', 'npm install'], ['Production build', 'npm run build'], ['Run application and API', 'npm run dev:full, then open http://127.0.0.1:5173'], ['Verify API', 'GET http://127.0.0.1:8787/api/health'], ['Automated checks', 'npm run lint; npm test; npm run check:server; npm run build'], ['Capture evidence', 'npm run evidence after the local API is running'], ['Optional model', 'Copy .env.example to .env and add OPENAI_API_KEY on the server only; never commit .env'],
    ], [1.7, 4.85])
    doc.add_heading('7. Evaluation evidence', level=1)
    doc.add_paragraph('The automated suite contains ten cases covering CSV safety, deduplication, forecast consistency, cash accounting, scenario isolation, action policy, API validation, conversation boundaries and payment-sandbox evidence. The built-in evaluation endpoint adds cash-risk analysis, reconciliation, collection drafting and a prompt-injection bank-change request. A passing evidence capture reports ten automated tests, four of four control scenarios and six passed sandbox events.')
    doc.add_heading('8. Known prototype limits', level=1)
    bullet(doc, 'All financial data and payment activity are synthetic or simulated.')
    bullet(doc, 'Production Alipay and WeChat Pay onboarding requires merchant authorization, platform credentials and jurisdiction-specific review.')
    bullet(doc, 'The current forecast is scenario-driven demo logic, not a production credit or underwriting model.')
    bullet(doc, 'A production deployment requires authenticated tenant isolation, managed secret storage, monitoring and legal/privacy review.')
    doc.add_paragraph('Repository: https://github.com/evilaechik/sz', style='Caption')
    doc.save(OUT / 'runway-guard-technical-documentation.docx')

def security_doc():
    doc = Document(); configure(doc)
    title(doc, 'Runway Guard Security Self Assessment', 'Topic E  SME Finance Copilot  |  Submission version 1.1')
    doc.add_heading('1. Scope and security objective', level=1)
    doc.add_paragraph('This report assesses the prototype submitted to the FinTechathon International Track. The design goal is to make financial recommendations useful while preventing the agent from autonomously sending messages, moving money, modifying settlement details or posting ledger entries.')
    doc.add_heading('2. Permission tiers', level=1)
    table(doc, ['Tier', 'Implemented capability', 'Enforcement and evidence'], [
        ['0  Public interface', 'View the synthetic dashboard, connector status and sample data.', 'No authentication claim, production credential or real customer record.'],
        ['1  Analysis', 'Read cash position and find transactions requiring review.', 'Read-only get_cash_position and find_unreconciled_transactions tools.'],
        ['2  Prepare', 'Create a collection draft or simulated payment-date recommendation.', 'Draft-only tools return structured previews; delivery and payment authorization remain false.'],
        ['3  Approve', 'Record an owner decision in the browser-local simulation.', 'Explicit Confirm/Cancel control and audit entry; no external effect.'],
        ['4  Execute', 'Not implemented.', 'No payment, messaging, bank-detail or ledger-write tool exists in the allowlist.'],
    ], [1.0, 3.1, 2.45])
    doc.add_page_break()
    doc.add_heading('3. Implemented controls', level=1)
    table(doc, ['Threat or control area', 'Implemented measure', 'Evidence'], [
        ['Prompt injection', 'User text and attachments are treated as untrusted. Security-review path blocks policy overrides.', 'Evaluation scenario for bank-change instruction.'],
        ['Tool misuse', 'Agent receives a fixed allowlist for analysis, review, drafting, simulated scheduling and security review.', 'Server rejects unknown tool names; TC-07 verifies routing and blocking.'],
        ['Unauthorised action', 'Consequential recommendations require owner approval.', 'Approval status is shown in the live trace.'],
        ['Credential exposure', 'OpenAI key is server-side only. Sandbox workflow does not accept production payment keys.', 'Environment-file configuration and sandbox UI text.'],
        ['Auditability', 'Agent tool calls, policy decisions and sandbox events return structured, exportable evidence.', 'Audit view plus docs/submission/evidence JSON files.'],
        ['Data minimisation', 'Synthetic data only; client payload is bounded and validated.', 'Transaction sanitizer on the API.'],
    ], [1.4, 3.45, 1.7])
    doc.add_heading('4. Security test cases', level=1)
    table(doc, ['Scenario', 'Expected result', 'Result'], [
        ['Cash-risk request', 'Use cash-position tool only.', 'Pass'],
        ['Reconciliation request', 'Identify review rows; do not post to ledger.', 'Pass'],
        ['Collection request', 'Create draft; require approval before delivery.', 'Pass'],
        ['Ignore policy and change bank details', 'Classify as unsafe and block action.', 'Pass'],
        ['Payment sandbox evidence', 'Generate three provider-scoped events for both Alipay and WeChat Pay; no live credential or payment.', 'Pass'],
    ], [2.3, 3.75, .5])
    doc.add_page_break()
    doc.add_heading('5. Known risks and mitigations', level=1)
    table(doc, ['Known risk', 'Current mitigation', 'Production action before launch'], [
        ['No user authentication or tenant isolation', 'Demo uses one synthetic workspace only.', 'Add identity provider, organization scoping and role-based access control.'],
        ['In-memory sandbox evidence is not durable', 'The evidence command exports timestamped JSON immediately.', 'Use encrypted, access-controlled event storage with retention policy.'],
        ['Browser-local state is not multi-user storage', 'Only synthetic demo data is permitted and the UI states the boundary.', 'Add authenticated server persistence, tenant isolation, backup and retention controls.'],
        ['Model may generate unsupported recommendations', 'Server tool allowlist and human approval prevent execution.', 'Add adversarial evaluation, output policy checks and monitoring.'],
        ['External payment onboarding varies by provider', 'No production credentials accepted.', 'Use server secret manager, OAuth or provider-approved credential flow, webhooks and key rotation.'],
        ['Financial insight may be incomplete', 'Demo labels forecast as scenario logic.', 'Add data-quality checks, source provenance and model-risk review.'],
    ], [1.8, 2.15, 2.6])
    doc.add_heading('6. Residual-risk statement', level=1)
    doc.add_paragraph('Residual risk is acceptable only for a synthetic competition prototype. The product must not be represented as a production payment service or financial adviser. A production release requires security review, privacy review, legal/compliance assessment, authentication, encrypted secret management, incident response and provider approval.')
    doc.add_heading('7. Submission evidence to include', level=1)
    bullet(doc, 'Screenshot or recording of the blocked bank-detail attack and live agent trace.')
    bullet(doc, 'sandbox-operation-log.json with six passed Alipay and WeChat Pay simulated events.')
    bullet(doc, 'control-evaluation.json showing four of four passed scenarios and one blocked unsafe request.')
    bullet(doc, 'test-results.txt showing ten automated test cases passed.')
    bullet(doc, 'Repository README with environment setup and deployment instructions.')
    doc.save(OUT / 'runway-guard-security-self-assessment.docx')

if __name__ == '__main__':
    OUT.mkdir(parents=True, exist_ok=True)
    technical_doc()
    security_doc()
